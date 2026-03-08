#!/usr/bin/env python3
"""Convert CSV tables to XLSX format."""
import csv
import os

try:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Installing openpyxl...")
    import subprocess
    subprocess.check_call(["pip", "install", "openpyxl"])
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

CSV_DIR = os.path.join(os.path.dirname(__file__), "..", "tables", "csv")
XLSX_DIR = os.path.join(os.path.dirname(__file__), "..", "tables", "xlsx")

os.makedirs(XLSX_DIR, exist_ok=True)

for filename in os.listdir(CSV_DIR):
    if not filename.endswith(".csv"):
        continue
    csv_path = os.path.join(CSV_DIR, filename)
    xlsx_path = os.path.join(XLSX_DIR, filename.replace(".csv", ".xlsx"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    with open(csv_path, "r", encoding="utf-8") as f:
        reader = csv.reader(f, delimiter=";")
        for row_idx, row in enumerate(reader, 1):
            for col_idx, value in enumerate(row, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)

    # Auto-adjust column widths
    for col_idx, column in enumerate(ws.columns, 1):
        max_length = max(len(str(cell.value or "")) for cell in column)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 2, 50)

    wb.save(xlsx_path)
    print(f"Created {xlsx_path}")
